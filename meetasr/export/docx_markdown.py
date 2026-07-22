"""Convert markdown-it tokens into python-docx document elements.

Visual decisions such as fonts, colours, spacing and page layout belong to the
DOCX template. This module maps Markdown semantics to named Word styles.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any, Sequence
from urllib.parse import urlsplit

from markdown_it.token import Token

from meetasr.export.markdown_parser import parse_markdown


_SAFE_HYPERLINK_SCHEMES = frozenset({"http", "https", "mailto"})


@dataclass
class _ListState:
    kind: str
    next_number: int = 1


def render_markdown(document: Any, markdown: str) -> None:
    """Append rendered Markdown blocks to a python-docx document/subdocument."""
    renderer = _DocxMarkdownRenderer(document)
    renderer.render(parse_markdown(markdown))


class _DocxMarkdownRenderer:
    def __init__(self, document: Any) -> None:
        self.document = document
        self.list_states: list[_ListState] = []
        self.item_paragraph_counts: list[int] = []
        self.quote_depth = 0

    def render(self, tokens: Sequence[Token]) -> None:
        index = 0
        while index < len(tokens):
            token = tokens[index]

            if token.type == "heading_open":
                level = min(max(int(token.tag[1:]), 1), 6)
                paragraph = self.document.add_paragraph(
                    style=self._style(f"Heading {level}", "Normal")
                )
                index = self._render_following_inline(paragraph, tokens, index)

            elif token.type == "paragraph_open":
                style, marker = self._paragraph_style_and_marker()
                paragraph = self.document.add_paragraph(style=style)
                if self.list_states and self.item_paragraph_counts:
                    self._apply_list_indent(paragraph)
                if marker:
                    paragraph.add_run(marker)
                index = self._render_following_inline(paragraph, tokens, index)
                if self.item_paragraph_counts:
                    self.item_paragraph_counts[-1] += 1

            elif token.type == "bullet_list_open":
                self.list_states.append(_ListState("bullet"))

            elif token.type == "ordered_list_open":
                start = token.attrGet("start")
                self.list_states.append(
                    _ListState("number", int(start) if start is not None else 1)
                )

            elif token.type in {"bullet_list_close", "ordered_list_close"}:
                if self.list_states:
                    self.list_states.pop()

            elif token.type == "list_item_open":
                self.item_paragraph_counts.append(0)

            elif token.type == "list_item_close":
                if self.item_paragraph_counts:
                    self.item_paragraph_counts.pop()

            elif token.type == "blockquote_open":
                self.quote_depth += 1

            elif token.type == "blockquote_close":
                self.quote_depth = max(0, self.quote_depth - 1)

            elif token.type == "table_open":
                index = self._render_table(tokens, index)

            elif token.type in {"fence", "code_block"}:
                self._render_code_block(token.content)

            elif token.type == "hr":
                paragraph = self.document.add_paragraph(style=self._style("Normal"))
                paragraph.add_run("―" * 24)

            index += 1

    def _render_following_inline(
        self,
        paragraph: Any,
        tokens: Sequence[Token],
        opening_index: int,
    ) -> int:
        inline_index = opening_index + 1
        if inline_index < len(tokens) and tokens[inline_index].type == "inline":
            _render_inline(paragraph, tokens[inline_index].children or [])
            return inline_index + 1
        return opening_index

    def _paragraph_style_and_marker(self) -> tuple[str, str | None]:
        if self.list_states and self.item_paragraph_counts:
            first_paragraph = self.item_paragraph_counts[-1] == 0
            if first_paragraph:
                state = self.list_states[-1]
                base = "List Bullet" if state.kind == "bullet" else "List Number"
                depth = min(len(self.list_states), 3)
                nested = base if depth == 1 else f"{base} {depth}"
                if self._has_style(nested):
                    if state.kind == "number":
                        state.next_number += 1
                    return nested, None
                if self._has_style(base):
                    if state.kind == "number":
                        state.next_number += 1
                    return base, None

                if state.kind == "bullet":
                    marker = "• "
                else:
                    marker = f"{state.next_number}. "
                    state.next_number += 1
                return self._style("Normal"), marker
        if self.quote_depth:
            return self._style("Quote", "Normal"), None
        return self._style("Normal"), None

    def _apply_list_indent(self, paragraph: Any) -> None:
        from docx.shared import Cm

        depth = len(self.list_states)
        paragraph.paragraph_format.left_indent = Cm(0.63 * depth)
        if self.item_paragraph_counts[-1] == 0:
            paragraph.paragraph_format.first_line_indent = Cm(-0.4)

    def _render_code_block(self, content: str) -> None:
        paragraph = self.document.add_paragraph(style=self._style("Quote", "Normal"))
        run = paragraph.add_run(content.rstrip("\n"))
        # Only the semantic monospace distinction is defined here; colours and
        # spacing remain controlled by the template's paragraph style.
        run.font.name = "Consolas"

    def _render_table(self, tokens: Sequence[Token], opening_index: int) -> int:
        rows: list[tuple[bool, list[list[Token]]]] = []
        row: list[list[Token]] | None = None
        cell: list[Token] | None = None
        in_header = False
        index = opening_index + 1

        while index < len(tokens) and tokens[index].type != "table_close":
            token = tokens[index]
            if token.type == "thead_open":
                in_header = True
            elif token.type == "thead_close":
                in_header = False
            elif token.type == "tr_open":
                row = []
            elif token.type in {"th_open", "td_open"}:
                cell = []
            elif token.type == "inline" and cell is not None:
                cell.extend(token.children or [])
            elif token.type in {"th_close", "td_close"} and row is not None:
                row.append(cell or [])
                cell = None
            elif token.type == "tr_close" and row is not None:
                rows.append((in_header, row))
                row = None
            index += 1

        column_count = max((len(cells) for _, cells in rows), default=0)
        if column_count == 0:
            return index

        table = self.document.add_table(rows=len(rows), cols=column_count)
        if self._has_style("Table Grid"):
            table.style = "Table Grid"

        for row_index, (is_header, cells) in enumerate(rows):
            for column_index, inline_tokens in enumerate(cells):
                paragraph = table.cell(row_index, column_index).paragraphs[0]
                _render_inline(paragraph, inline_tokens, force_bold=is_header)
        return index

    def _has_style(self, name: str) -> bool:
        try:
            self.document.styles[name]
        except KeyError:
            return False
        return True

    def _style(self, *candidates: str) -> str:
        for candidate in candidates:
            if self._has_style(candidate):
                return candidate
        return "Normal"


def _render_inline(
    paragraph: Any,
    tokens: Sequence[Token],
    *,
    force_bold: bool = False,
) -> None:
    bold_depth = 0
    italic_depth = 0
    strike_depth = 0
    index = 0

    while index < len(tokens):
        token = tokens[index]

        if token.type == "strong_open":
            bold_depth += 1
        elif token.type == "strong_close":
            bold_depth = max(0, bold_depth - 1)
        elif token.type == "em_open":
            italic_depth += 1
        elif token.type == "em_close":
            italic_depth = max(0, italic_depth - 1)
        elif token.type == "s_open":
            strike_depth += 1
        elif token.type == "s_close":
            strike_depth = max(0, strike_depth - 1)
        elif token.type == "link_open":
            closing_index = _matching_close(tokens, index, "link_open", "link_close")
            label = _inline_plain_text(tokens[index + 1 : closing_index])
            href = token.attrGet("href") or ""
            if _is_safe_hyperlink(href):
                _add_hyperlink(
                    paragraph,
                    label or href,
                    href,
                    bold=force_bold or bold_depth > 0,
                    italic=italic_depth > 0,
                    strike=strike_depth > 0,
                )
            else:
                run = paragraph.add_run(label or href)
                _format_run(
                    run,
                    bold=force_bold or bold_depth > 0,
                    italic=italic_depth > 0,
                    strike=strike_depth > 0,
                )
            index = closing_index
        elif token.type == "image":
            alt_text = token.content or _inline_plain_text(token.children or [])
            run = paragraph.add_run(f"[Hình: {alt_text or 'ảnh'}]")
            _format_run(
                run,
                bold=force_bold or bold_depth > 0,
                italic=italic_depth > 0,
                strike=strike_depth > 0,
            )
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif token.type == "code_inline":
            run = paragraph.add_run(token.content)
            _format_run(
                run,
                bold=force_bold or bold_depth > 0,
                italic=italic_depth > 0,
                strike=strike_depth > 0,
            )
            run.font.name = "Consolas"
        elif token.type in {"text", "html_inline"}:
            run = paragraph.add_run(token.content)
            _format_run(
                run,
                bold=force_bold or bold_depth > 0,
                italic=italic_depth > 0,
                strike=strike_depth > 0,
            )
        index += 1


def _format_run(run: Any, *, bold: bool, italic: bool, strike: bool) -> None:
    run.bold = bold
    run.italic = italic
    run.font.strike = strike


def _matching_close(
    tokens: Sequence[Token],
    opening_index: int,
    opening_type: str,
    closing_type: str,
) -> int:
    depth = 0
    for index in range(opening_index, len(tokens)):
        if tokens[index].type == opening_type:
            depth += 1
        elif tokens[index].type == closing_type:
            depth -= 1
            if depth == 0:
                return index
    return opening_index


def _inline_plain_text(tokens: Sequence[Token]) -> str:
    parts: list[str] = []
    for token in tokens:
        if token.type in {"text", "code_inline", "html_inline"}:
            parts.append(token.content)
        elif token.type == "image":
            parts.append(token.content or "ảnh")
        elif token.type in {"softbreak", "hardbreak"}:
            parts.append("\n")
    return "".join(parts)


def _is_safe_hyperlink(url: str) -> bool:
    if not url or any(character in url for character in "\r\n\x00"):
        return False
    return urlsplit(url).scheme.lower() in _SAFE_HYPERLINK_SCHEMES


def _add_hyperlink(
    paragraph: Any,
    text: str,
    url: str,
    *,
    bold: bool,
    italic: bool,
    strike: bool,
) -> None:
    """Append an external hyperlink without fetching its target."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    relationship_id = paragraph.part.relate_to(
        url,
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = paragraph.add_run(text)
    try:
        run.style = paragraph.part.document.styles["Hyperlink"]
    except KeyError:
        styles = paragraph.part.document.styles
        hyperlink_style = styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
        hyperlink_style.font.color.rgb = _rgb_color("0563C1")
        hyperlink_style.font.underline = True
        run.style = hyperlink_style
    _format_run(run, bold=bold, italic=italic, strike=strike)

    paragraph._p.remove(run._r)
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def _rgb_color(value: str) -> Any:
    from docx.shared import RGBColor

    return RGBColor.from_string(value)
